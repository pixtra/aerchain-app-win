"""Engine + pipeline + aerbot + extraction golden tests.

Run:  .venv/bin/pytest -q
"""
from __future__ import annotations

import os
import sys

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app import pipeline, seed_data  # noqa: E402
from app.database import get_conn, init_schema  # noqa: E402
from app.domain.enums import (  # noqa: E402
    Confidence,
    CoverageStatus,
    MappingStatus,
    ResponseStatus,
)
from app.engines.discount import (  # noqa: E402
    DiscountResult,
    applicable_tier,
    calculate_discount,
)
from app.engines.economics import effective_cost  # noqa: E402
from app.engines.financing import financing_benefit  # noqa: E402
from app.engines.freight import allocate_freight  # noqa: E402
from app.engines.fx import apply_fx  # noqa: E402
from app.engines.uom import normalize_uom  # noqa: E402


# ------------------------------------------------------------------- fixtures
@pytest.fixture(scope="module")
def dataset():
    return seed_data.build()


@pytest.fixture(scope="module")
def conn(dataset):
    pipeline.reset_db(dataset)
    c = get_conn()
    yield c
    c.close()


@pytest.fixture(autouse=True)
def _isolate_conditions():
    """Award tests assume a clean baseline, but buyer conditions are live user
    data that survives re-seed by design. Deactivate around each test, restore
    after — never delete what isn't ours."""
    from app.services import conditions as condsvc
    prior = [(c["condition_id"], c["active"]) for c in condsvc.list_conditions()]
    for cid, _ in prior:
        condsvc.set_active(cid, False)
    yield
    current = {c["condition_id"] for c in condsvc.list_conditions()}
    for cid, active in prior:
        if cid in current:
            condsvc.set_active(cid, active)


# ------------------------------------------------------------- golden dataset
def test_golden_counts(dataset):
    assert len(dataset["skus"]) == 50
    assert len(dataset["suppliers"]) == 8
    assert len(dataset["products"]) == 249
    assert len(dataset["lines"]) == 30
    assert len(dataset["offers"]) == 131
    assert len(dataset["tiers"]) == 126
    assert len(dataset["evidence"]) == 131
    assert len(dataset["q_responses"]) == 40


def test_golden_edge_cases(dataset):
    sup8 = next(s for s in dataset["rfx_suppliers"] if s.supplier_id == "SUP-008")
    assert sup8.response_status == ResponseStatus.NO_RESPONSE
    ever = {o.supplier_id for o in dataset["offers"]}
    assert "SUP-008" not in ever
    l29_30 = {l.sku_id for l in dataset["lines"] if l.line_number >= 29}
    # L29/L30 are mapped to supplier products but nobody quoted them
    mapped = {p.sku_id for p in dataset["products"]
              if p.mapping_status != MappingStatus.UNRESOLVED}
    assert l29_30 <= mapped
    quoted_lines = {o.line_number for o in dataset["offers"]}
    assert not ({29, 30} & quoted_lines)
    unres_prod = {p.sku_id for p in dataset["products"]
                  if p.mapping_status == MappingStatus.UNRESOLVED}
    assert unres_prod  # unresolved mappings exist and are surfaced
    unresolved = [o for o in dataset["offers"]
                  if o.supplier_id == "SUP-004" and o.quoted_uom == "UNKNOWN"]
    assert len(unresolved) == 2


def test_golden_currency_and_expiry(dataset):
    usd = [o for o in dataset["offers"] if o.supplier_id == "SUP-003"]
    assert usd and all(o.quoted_currency == "USD" for o in usd)
    from datetime import date
    assert all(date.fromisoformat(o.quote_valid_until)
               > date.fromisoformat(o.quote_received_date)
               for o in dataset["offers"])


# ------------------------------------------------------------------ engines
def test_fx_identity_and_conversion(dataset):
    rates = {"USD": {"INR": 86.50}}
    assert apply_fx(10.0, "INR", "INR", {}).normalized == 10.0
    r = apply_fx(1.0, "USD", "INR", rates)
    assert r.normalized == 86.50 and r.fx_rate_id == "FX-USD-INR"
    fx = {x.from_currency: x.fx_rate for x in dataset["fx"]}
    assert fx["USD"] == 86.50 and fx["EUR"] == 93.20 and fx["GBP"] == 108.60


def test_uom_normalize():
    r = normalize_uom(quoted_unit_price=800.0, conversion_factor=2.0, source="MASTER")
    assert r.normalized_unit_price == pytest.approx(400.0)
    assert r.conversion_source == "MASTER" and not r.conflict


def test_discount_tier_lookup(dataset):
    from app.domain.models import DiscountTier
    tiers = [
        DiscountTier("DT-T1", "OFF-0001", 1, 0.0, 99.0, "UNIT",
                     discount_percentage=1.5, normalized_min_qty=0.0,
                     normalized_max_qty=100.0),
        DiscountTier("DT-T2", "OFF-0001", 2, 100.0, 499.0, "UNIT",
                     discount_percentage=3.0, normalized_min_qty=100.0,
                     normalized_max_qty=None),
    ]
    picked = applicable_tier(250.0, tiers, "EA")
    assert picked["pct"] == 3.0
    res = calculate_discount(100.0, 250.0, "FLAT_PERCENT", 5.0, None, None)
    assert isinstance(res, DiscountResult)
    assert res.benefit == pytest.approx(5.0)


def test_effective_cost_baseline(dataset):
    assert effective_cost(100.0, 5.0, 1.0, 0.5) == pytest.approx(95.5)


def test_financing_and_freight(dataset):
    fin = financing_benefit(100.0, 60, 0.12, 30)
    assert fin == pytest.approx(100.0 * 0.12 * 30 / 365.0)
    fr = allocate_freight("EXTRA_PER_SHIPMENT", 1000.0, 5000.0)
    assert fr.allocated_per_unit == pytest.approx(0.2)
    assert allocate_freight("INCLUDED", 0.0, 5000.0).allocated_per_unit == 0


# ------------------------------------------------------------------ pipeline
def test_pipeline_loads(conn):
    def one(q):
        return conn.execute(q).fetchone()[0]
    assert one("SELECT COUNT(*) FROM supplier_offer") == 131
    assert one("SELECT COUNT(*) FROM decision_ready") == 131
    assert one("SELECT COUNT(*) FROM coverage") == 240
    assert one("SELECT COUNT(*) FROM award_decision") == 28
    el = one("SELECT COUNT(*) FROM decision_ready WHERE overall_eligible=1")
    assert 45 <= el <= 110


def test_pipeline_eligible_balance(conn):
    by_sup = dict(conn.execute(
        "SELECT supplier_id, SUM(overall_eligible) FROM decision_ready "
        "GROUP BY supplier_id").fetchall())
    assert by_sup.get("SUP-006", 0) == 0 and by_sup.get("SUP-007", 0) == 0
    assert by_sup.get("SUP-001", 0) > by_sup.get("SUP-007", 0) + 10


def test_pipeline_reason_text(conn):
    reasons = conn.execute("SELECT eligibility_reason FROM decision_ready "
                           "WHERE overall_eligible=0 AND eligibility_reason!='' "
                           "LIMIT 1").fetchone()
    assert reasons


def test_pipeline_no_hardcoded_values(conn):
    prices = {r[0] for r in conn.execute(
        "SELECT DISTINCT normalized_unit_price FROM decision_ready")}
    assert len(prices) > 100  # 131 offers must be individually computed


# ---------------------------------------------------------------- aerbot
class ScriptedLLM:
    """Simulates a real model's tool-calling behaviour: picks a tool from the
    question, reads the serialized tool output, and produces a grounded final
    JSON envelope whose figures are copied verbatim from that output."""

    def __init__(self):
        self.tool_calls = []
        self.tool_payloads = []
        self.copied_value = None

    def chat_fn(self, system, messages, tools):
        from app.aerbot.llm import ChatResponse
        question = next(m["content"] for m in messages if m["role"] == "user")
        if any(m["role"] == "tool" for m in messages):
            return ChatResponse(content=self._final(messages, question))
        calls = self._plan(question)
        self.tool_calls.append(calls)
        return ChatResponse(tool_calls=calls)

    def _plan(self, q):
        from app.aerbot.llm import ToolCall
        low = q.lower()
        if "cheapest" in low:
            return [ToolCall(id="t1", name="get_line_comparison",
                             arguments={"line_number": 3})]
        if "award" in low:
            return [ToolCall(id="t1", name="run_award_scenario",
                             arguments={"scenario_id": "SCN-T", "scenario_name": "Test"})]
        if "coverage" in low:
            return [ToolCall(id="t1", name="get_coverage", arguments={})]
        if "uncertain" in low:
            return [ToolCall(id="t1", name="get_uncertainties", arguments={})]
        if "what if" in low or "scenario" in low:
            return [ToolCall(id="t1", name="run_cost_scenario",
                             arguments={"scenario_name": "What-if",
                                        "additional_discounts": {"SUP-002": 5.0}})]
        if "calculate" in low:
            from app.aerbot.tools import get_line_comparison
            off = next(r["offer_id"] for r in get_line_comparison(line_number=12)
                       if r["supplier_id"] == "SUP-002")
            return [ToolCall(id="t1", name="get_offer_calculation",
                             arguments={"offer_id": off}),
                    ToolCall(id="t2", name="get_offer_evidence",
                             arguments={"offer_id": off})]
        return [ToolCall(id="t1", name="get_rfx_summary", arguments={})]

    def _final(self, messages, question):
        import json
        payloads = [json.loads(m["content"]) for m in messages if m["role"] == "tool"]
        self.tool_payloads = payloads
        node = payloads[0] if payloads else None
        self.copied_value = self._pick(node) if node else "N/A"
        summary = ""
        if isinstance(node, dict) and node.get("summary"):
            summary = node["summary"] + " "
        return json.dumps({
            "intent": "comparison", "title": "Scripted-model answer",
            "narrative": (f"{summary}Executive read: the computed figure is "
                          f"{self.copied_value}, copied verbatim."),
            "table": [{"value": self.copied_value}],
            "assumptions": ["Figures come from deterministic engines"],
            "exclusions": [], "uncertainties": [],
            "evidence": [{"offer_id": "FAKE-999", "source_file": "invented.pdf"}],
        }, ensure_ascii=False)

    def _pick(self, node):
        if isinstance(node, list) and node and isinstance(node[0], dict):
            return node[0].get("effective_economic_cost") or node[0].get("offer_count")
        if isinstance(node, dict):
            for k in ("eligible_count", "line_count", "effective_economic_cost"):
                if node.get(k):
                    return node[k]
            if node.get("awards"):
                return node["awards"][0]["offer_id"]
        return "N/A"


@pytest.fixture
def scripted():
    return ScriptedLLM()


def test_aerbot_agent_grounds_numbers(conn, scripted):
    """The model's numbers must equal the engines' exact serialized output."""
    from app.aerbot import answer
    from app.aerbot.tools import get_line_comparison
    a = answer("Who is the cheapest supplier on line 3?", chat_fn=scripted.chat_fn)
    assert a["title"] and a["narrative"]
    # the scripted model copied the real tool's serialized money string verbatim
    real = serialize_money(get_line_comparison(line_number=3)[0])
    assert scripted.copied_value and scripted.copied_value != "N/A"
    assert scripted.copied_value == real, "narrative figure must match engine output"
    assert scripted.copied_value in a["narrative"]
    # invented evidence must not survive grounding
    assert all(e.get("source_file") != "invented.pdf" for e in a["evidence"])


def test_aerbot_all_intents_route(conn, scripted):
    from app.aerbot import answer
    cases = ["Give me a summary of the RFx",
             "Who is the cheapest supplier on line 3?",
             "Show me the recommended award split",
             "Which lines have poor coverage?",
             "What are we uncertain about?",
             "How did you calculate supplier B's effective cost on line 12?",
             "What if supplier B offered 5% additional discount?"]
    for q in cases:
        scripted.__init__()
        a = answer(q, chat_fn=scripted.chat_fn)
        assert a["title"] and a["narrative"]
        assert scripted.tool_payloads, f"{q!r} executed no tools"
        assert scripted.tool_calls, "agent loop must invoke the model"


def test_aerbot_evidence_real_files(conn, scripted):
    from app.aerbot.tools import get_line_comparison
    from app.aerbot import answer
    q = "How did you calculate supplier B's effective cost on line 12?"
    a = answer(q, chat_fn=scripted.chat_fn)
    real = {r[0] for r in conn.execute("SELECT DISTINCT source_file FROM offer_evidence")}
    if a["evidence"]:
        assert all(e["source_file"] in real for e in a["evidence"])


def test_aerbot_status_informs_config(monkeypatch):
    from app.aerbot import status
    for k in ("OPENAI_API_KEY", "OPENAI_API_BASE", "KTQ_LLM_URL", "KTQ_LLM_KEY",
              "ANTHROPIC_API_KEY", "OLLAMA_HOST"):
        monkeypatch.delenv(k, raising=False)
    st = status()
    if not st["configured"]:
        assert "OPENAI_API_KEY" in st["detail"] or "Ollama" in st["detail"]


def test_aerbot_rejects_invented_numbers(conn, scripted):
    """A model that cites a ₹ figure never seen in tool output must be
    forced to correct it; the invented figure never reaches the answer."""
    from app.aerbot import answer
    from app.aerbot.tools import get_line_comparison

    real = serialize_money(get_line_comparison(line_number=3)[0])
    original_final = scripted._final
    first = {"count": 0}

    def bad_then_good(messages, question):
        import json
        if first["count"] == 0:
            first["count"] += 1
            return json.dumps({
                "intent": "comparison", "title": "Bad",
                "narrative": f"cost is ₹ 9999.9999 and the real figure is {real}",
                "table": [], "assumptions": [], "exclusions": [],
                "uncertainties": [], "evidence": []}, ensure_ascii=False)
        return original_final(messages, question)

    scripted._final = bad_then_good
    a = answer("Who is the cheapest eligible supplier on line 3?",
               chat_fn=scripted.chat_fn)
    assert "9999.9999" not in a["narrative"]
    assert real in a["narrative"] or "₹ 393.6464" in a["narrative"]
    assert first["count"] == 1, "grounding guard must have forced a rewrite"


def test_aerbot_award_summary_is_verbatim(conn, scripted):
    """The deterministic award summary must appear word-for-word in the answer."""
    from app.aerbot import answer
    from app.aerbot.tools import run_award_scenario, make_award_scenario
    real_summary = run_award_scenario(make_award_scenario("SCN-T", "test"))["summary"]
    a = answer("Show me the recommended award split", chat_fn=scripted.chat_fn)
    assert real_summary in a["narrative"], "engine summary must be reproduced verbatim"
    # and the tool really reports per-supplier counts (not "1 each")
    assert "SUP-001 10" in real_summary


def test_parse_envelope_salvages_truncated_json():
    """Small local models sometimes truncate the envelope; the narrative must
    still be recovered instead of looping until the step budget is spent."""
    from app.aerbot.agent import _parse_envelope
    truncated = ('{"intent":"summary","title":"T","narrative":"All good, 131 offers",'
                 '"table":[{"value":1')
    env = _parse_envelope(truncated)
    assert env is not None and env["narrative"] == "All good, 131 offers"


def test_aerbot_requires_a_tool_call(conn):
    """An answer produced without consulting any engine must be rejected."""
    import json
    from app.aerbot import answer
    from app.aerbot.llm import ChatResponse, ToolCall

    state = {"n": 0}

    def chat_fn(system, messages, tools):
        state["n"] += 1
        if any(m["role"] == "tool" for m in messages):
            return ChatResponse(content=json.dumps({
                "intent": "summary", "title": "Grounded",
                "narrative": "The RFx has eligible offers.",
                "table": [], "assumptions": [], "exclusions": [],
                "uncertainties": [], "evidence": []}))
        if state["n"] == 1:
            return ChatResponse(content=json.dumps({
                "intent": "summary", "title": "Unchecked",
                "narrative": "The RFx has 999 eligible offers.",
                "table": [], "assumptions": [], "exclusions": [],
                "uncertainties": [], "evidence": []}))
        return ChatResponse(tool_calls=[ToolCall(id="t1", name="get_rfx_summary",
                                                 arguments={})])

    a = answer("Give me a summary of the RFx", chat_fn=chat_fn)
    assert state["n"] >= 3, "a no-tool answer must be rejected and a tool forced"
    assert "999" not in a["narrative"]


def test_aerbot_prepends_engine_summary(conn, scripted):
    """If the model omits the engine summary, the server prepends it verbatim
    rather than looping the model until the budget dies."""
    import json
    from app.aerbot import answer
    from app.aerbot.tools import run_award_scenario, make_award_scenario
    real_summary = run_award_scenario(make_award_scenario("SCN-T", "test"))["summary"]
    original = scripted._final

    def no_summary(messages, question):
        env = json.loads(original(messages, question))
        env["narrative"] = "SUP-001 leads this split based on lowest effective cost."
        return json.dumps(env)

    scripted._final = no_summary
    a = answer("Show me the recommended award split", chat_fn=scripted.chat_fn)
    assert real_summary in a["narrative"]


def serialize_money(row):
    import json
    from app.aerbot.domain_serialize import serialize_tool_result
    d = json.loads(serialize_tool_result([row]))
    return d[0]["effective_economic_cost"]


# ---------------------------------------------------------------- extraction
import shutil as _shutil

_needs_tess = pytest.mark.skipif(not _shutil.which("tesseract"),
                                 reason="tesseract not installed")
_needs_pdftoppm = pytest.mark.skipif(not _shutil.which("pdftoppm"),
                                     reason="poppler not installed")

@_needs_tess
def test_extraction_roundtrip(dataset):
    from app.services.extraction import extract_all
    os.makedirs(os.path.join(os.path.dirname(__file__), "../data/raw"),
                exist_ok=True)
    from scripts.generate_artifacts import main as gen
    gen()
    res = extract_all()
    sup_of = {"Supplier_A_Quote.xlsx": "SUP-001", "Supplier_B_Quote.pdf": "SUP-002",
              "Supplier_C_Quote.xlsx": "SUP-003", "Supplier_D_Rate_Card.jpg": "SUP-004",
              "Supplier_E_Email.txt": "SUP-005", "Supplier_F_Quote.pdf": "SUP-006",
              "Supplier_G_Quote.xlsx": "SUP-007"}
    by_sup = {s.supplier_id: s for s in dataset["suppliers"]}
    for fn, sup in sup_of.items():
        if fn == "Supplier_D_Rate_Card.jpg":
            continue  # image goes through real OCR below (fuzzy by nature)
        rows = res[fn]["rows"]
        assert len(rows) == sum(1 for o in dataset["offers"]
                                if o.supplier_id == sup), fn
    img = res["Supplier_D_Rate_Card.jpg"]["rows"]
    assert all(r.get("ocr_engine") == "tesseract" for r in img), \
        "image path must use the real OCR engine, never the sidecar"
    assert len(img) >= 12, "real OCR must recover most rate-card rows"
    assert sum(1 for r in img if r.get("unresolved")) >= 1, \
        "genuine OCR misreads must surface as unresolved, not silently pass"


def test_supplier008_no_response(dataset):
    sup8 = next(s for s in dataset["rfx_suppliers"] if s.supplier_id == "SUP-008")
    assert sup8.response_status == ResponseStatus.NO_RESPONSE


@_needs_pdftoppm
@_needs_tess
def test_scanned_pdf_routes_to_ocr(tmp_path):
    """An image-only PDF has no text layer — it must be rendered + OCR'd,
    not silently return zero rows."""
    from PIL import Image
    from app.services.extraction import extract_pdf_text, pdf_source_engine
    src = os.path.join(os.path.dirname(__file__), "../data/raw/Supplier_D_Rate_Card.jpg")
    pdf = str(tmp_path / "scan.pdf")
    Image.open(src).save(pdf, "PDF")
    assert pdf_source_engine(pdf) == "tesseract-scanned-pdf"
    rows = extract_pdf_text(pdf)
    assert len(rows) >= 10
    assert sum(1 for r in rows if r.get("unresolved")) >= 1


def test_tier_hints_parse():
    from app.services.extraction import parse_tier_hints
    rules = parse_tier_hints(["5% above 500 units", "1000+: 7.5%",
                              "no discount here", "110% above 10"])
    assert [(r["discount_percentage"], r["supplier_min_qty"]) for r in rules] == [
        (5.0, 500.0), (7.5, 1000.0)]


def test_promote_upload_end_to_end(conn):
    """Promoted rows run the real engines and land in decision_ready/award;
    incomplete rows are skipped with reasons, and the award is restored after."""
    from types import SimpleNamespace
    from app.domain.models import Scenario
    from app.engines.scenario import award_line_level
    from app.pipeline import _insert, _to_row
    from app.services import ingest

    code = next(r[0] for r in conn.execute(
        "SELECT supplier_product_code FROM supplier_product"
        " WHERE supplier_id='SUP-001' LIMIT 1"))
    rows = ingest.match_rows([{
        "sku_code": code, "quantity": 100.0, "unit_uom": "BOX",
        "currency": "INR", "unit_price": 10.0, "moq": 0.0,
        "lead_days": 10, "valid_until": "2026-12-31",
        "unresolved": [], "raw_text": "5% above 500 units"}])
    assert rows[0]["matched_supplier_id"] == "SUP-001"
    fid = ingest.store_upload("Promo.txt", "/tmp/promo.txt", "EMAIL", rows,
                              "regex-text", b"promo")["file_id"]
    try:
        # supplier scope: another supplier's catalogue must NOT match
        scoped = ingest.match_rows([dict(rows[0])], supplier_id="SUP-002")
        assert scoped[0]["matched_supplier_id"] is None
        assert "no_product_match" in scoped[0]["unresolved"]
        # duplicate content rejected
        try:
            ingest.store_upload("Promo.txt", "/tmp/promo.txt", "EMAIL", rows,
                                "regex-text", b"promo")
            assert False, "duplicate upload must be rejected"
        except ValueError:
            pass
        res = ingest.promote_upload(fid)
        assert res["promoted"] >= 1 and res["tiers_applied"] == 1
        n = conn.execute("SELECT COUNT(*) FROM decision_ready"
                         " WHERE offer_id LIKE 'OFF-U%'").fetchone()[0]
        assert n == res["promoted"]
        total = conn.execute("SELECT COUNT(*) FROM decision_ready").fetchone()[0]
        assert total == 131 + res["promoted"], \
            "promote must append, never overwrite golden rows"
    finally:
        for t, col in (("supplier_offer", "offer_id"),
                       ("discount_tier", "offer_id"),
                       ("offer_evidence", "offer_id"),
                       ("decision_ready", "offer_id")):
            conn.execute(f"DELETE FROM {t} WHERE {col} LIKE 'OFF-U%'")
        conn.execute("DELETE FROM uploaded_offer WHERE file_id=?", (fid,))
        conn.execute("DELETE FROM uploaded_file WHERE file_id=?", (fid,))
        recs = [SimpleNamespace(**dict(r)) for r in
                conn.execute("SELECT * FROM decision_ready")]
        pol = SimpleNamespace(**dict(
            conn.execute("SELECT * FROM policy LIMIT 1").fetchone()))
        base = Scenario(scenario_id="BASE", rfx_id="RFX-001",
                        scenario_name="Baseline award · cheapest eligible per line",
                        award_strategy="LINE_LEVEL_SPLIT")
        conn.execute("DELETE FROM award_decision")
        for a in award_line_level(recs, base, pol):
            _insert(conn, "award_decision", _to_row(a), conjoin=True)
        conn.commit()
        assert conn.execute("SELECT COUNT(*) FROM award_decision").fetchone()[0] == 28


# --------------------------------------------------------------------- ingest
def test_canon_normalizes_nl_variants():
    """'Quantity->500 BOX 120 rupees' (jpg) and 'Qty: 500 box Rs. 120' (pdf)
    must parse to the same fact — dialects fold before positional parsing."""
    from app.services.extraction import _parse_keyvalue_rows, _parse_ocr_rows
    jpg_style = "D-003 Quantity->500 BOX 120 rupees"
    pdf_style = "D-003 Qty: 500 box Rs. 120"
    ra = _parse_ocr_rows([jpg_style])[0]
    rb = _parse_ocr_rows([pdf_style])[0]
    assert (ra["quantity"], ra["currency"], ra["unit_uom"]) == (500, "INR", "BOX")
    assert (rb["quantity"], rb["currency"], rb["unit_uom"]) == (500, "INR", "BOX")
    assert ra["unit_price"] == rb["unit_price"] == 120
    rc = _parse_keyvalue_rows(["D-003 Qty: 500 box Rs. 120 10d valid 2026-11-30"])[0]
    assert (rc["quantity"], rc["currency"]) == (500, "INR")


@_needs_tess
def test_card_layout_parses(conn):
    """Phone-screenshot product cards (facts scattered across lines) parse
    into one honest row instead of zero rows."""
    from app.services.extraction import parse_card_blocks
    from app.services.ingest import match_rows
    text = ("Supplier Name : EcoPack Solutions |\n"
            "Quoted: INR 0.01/BAG (5% above 500 INR)\n\n"
            "SKU Item 0036\n\nQuantity = 20000")
    rows = parse_card_blocks(text)
    assert len(rows) == 1
    r = rows[0]
    assert (r["quantity"], r["unit_uom"], r["currency"], r["unit_price"]) == \
        (20000.0, "BAG", "INR", 0.01)
    assert r["supplier_name_hint"] == "EcoPack Solutions"
    m = match_rows(rows)[0]
    assert (m["matched_supplier_id"], m["matched_sku_id"]) == ("SUP-005", "Item-036")


def test_ingest_upload_real_ocr_and_match(conn):
    """A user-supplied image goes through real Tesseract OCR and rows match
    the product master — no sidecar, no invented values."""
    from app.services import ingest
    src = os.path.join(os.path.dirname(__file__), "../data/raw/Supplier_D_Rate_Card.jpg")
    with open(src, "rb") as fh:
        path = ingest.save_upload("TestCard.jpg", fh.read())
    try:
        rows, engine = ingest.extract_upload(path, "IMAGE")
        assert engine == "tesseract"
        assert len(rows) >= 12
        matched = ingest.match_rows(rows)
        assert sum(1 for r in matched if r.get("matched_supplier_id")) >= 10
        summary = ingest.store_upload("TestCard.jpg", path, "IMAGE", matched, engine)
        assert summary["rows"] == len(rows)
        assert len(ingest.get_upload_rows(summary["file_id"])) == len(rows)
    finally:
        conn.execute("DELETE FROM uploaded_offer WHERE file_id IN "
                     "(SELECT file_id FROM uploaded_file WHERE filename='TestCard.jpg')")
        conn.execute("DELETE FROM uploaded_file WHERE filename='TestCard.jpg'")
        conn.commit()
        if os.path.exists(path):
            os.remove(path)

# ---------------------------------------------------------------- conditions
def test_condition_nl_parses_all_kinds():
    from app.services.conditions import parse_condition
    cases = [
        ("exclude SUP-004", "EXCLUDE_SUPPLIER", "SUP-004", None),
        ("no supplier more than 40%", "MAX_SUPPLIER_SHARE", None, 40.0),
        ("at least 3 suppliers", "MIN_SUPPLIERS", None, 3),
        ("lead time max 10 days", "MAX_LEAD_DAYS", None, 10.0),
        ("require quality pass", "REQUIRE_QUALITY_PASS", None, None),
        ("quotes valid until 2026-12-31", "REQUIRE_VALID_UNTIL", None, None),
        ("nothing above 500 per unit", "MAX_EFFECTIVE_COST", None, 500.0),
    ]
    for text, kind, target, num in cases:
        cond, err = parse_condition(text)
        assert err is None, text
        assert cond["kind"] == kind, text
        assert cond.get("target") == target, text
        assert (cond.get("value_num") == num if num is not None
                else cond.get("value_num") is None), text
    cond, err = parse_condition("make it cheap please")
    assert cond is None and "could not understand" in err


def test_conditions_enforced_at_award(conn):
    from app.services import conditions as condsvc
    from app.aerbot.tools import make_award_scenario, run_award_scenario
    base = run_award_scenario(make_award_scenario("SCN-T", "t"))
    assert base["line_count"] == 28
    cid = condsvc.add_condition("exclude SUP-001", source="test")["condition_id"]
    try:
        out = run_award_scenario(make_award_scenario("SCN-T", "t"))
        assert "SUP-001" not in out["awards_by_supplier"]
        assert out["line_count"] < 28
        assert any("exclude SUP-001" in s for s in out["conditions_applied"])
    finally:
        condsvc.delete_condition(cid)
    assert run_award_scenario(make_award_scenario("SCN-T", "t"))["line_count"] == 28


def test_aerbot_condition_tools_route(conn, scripted):
    from app.aerbot import answer
    from app.services import conditions as condsvc
    from app.aerbot.llm import ChatResponse, ToolCall
    calls = {"n": 0}

    def chat_fn(system, messages, tools):
        names = [t["name"] for t in tools]
        assert "set_buyer_condition" in names and "get_buyer_conditions" in names
        calls["n"] += 1
        if calls["n"] == 1:
            return ChatResponse(tool_calls=[ToolCall(
                id="t1", name="set_buyer_condition",
                arguments={"text": "exclude SUP-007"})])
        import json
        return ChatResponse(content=json.dumps({
            "intent": "award", "title": "Condition set",
            "narrative": "SUP-007 is now excluded from the award.",
            "table": [], "assumptions": [], "exclusions": [],
            "uncertainties": [], "evidence": []}))

    a = answer("exclude SUP-007 from the award", chat_fn=chat_fn)
    assert "SUP-007" in a["narrative"]
    conds = [c for c in condsvc.list_conditions() if c["target"] == "SUP-007"]
    assert conds, "aerbot must really store the condition"
    for c in conds:
        condsvc.delete_condition(c["condition_id"])


# ------------------------------------------------------- flexible conditions
def test_region_mapping():
    from app.services.conditions import supplier_country, supplier_region
    assert supplier_region("Pune, IN") == "Asia"
    assert supplier_region("Singapore") == "Asia"
    assert supplier_region("Vietnam") == "Asia"
    assert supplier_country("Pune, IN") == "IN"


def test_v2_deterministic_specs():
    from app.services.conditions import parse_v2
    spec, _ = parse_v2("exclude suppliers outside Asia")
    assert spec["effect"] == "exclude" and spec["all"][0]["value"] == "Asia"
    spec, _ = parse_v2("only suppliers from India")
    assert spec["all"][0] == {"attr": "country", "op": "==", "value": "IN"}
    spec, _ = parse_v2("only quantity-matching suppliers for Item-036")
    assert spec["scope"] == {"sku": "Item-036"}
    assert spec["all"][0]["attr"] == "quantity_match"
    spec, _ = parse_v2("payment under Net 45")
    assert spec["all"][0] == {"attr": "payment_days", "op": "<=", "value": 45}


def test_model_fallback_validated():
    import json
    from app.services.conditions import model_parse_condition
    from app.aerbot.llm import ChatResponse

    def fake_chat(system, messages, tools):
        assert "SUP-001" in system  # live masters in the prompt
        return ChatResponse(content=json.dumps({
            "effect": "exclude", "entity": "supplier", "scope": None,
            "all": [{"attr": "region", "op": "!=", "value": "Asia"}]}))

    spec, desc = model_parse_condition("no Asians", chat_fn=fake_chat)
    assert spec["all"][0]["value"] == "Asia"

    def fake_bad(system, messages, tools):
        return ChatResponse(content=json.dumps({
            "effect": "exclude", "entity": "supplier", "scope": None,
            "all": [{"attr": "planet", "op": "==", "value": "Mars"}]}))

    try:
        model_parse_condition("no Martians", chat_fn=fake_bad)
        assert False, "unknown attrs must be rejected"
    except ValueError:
        pass


def test_custom_rule_enforced_at_award(conn):
    from app.services import conditions as condsvc
    from app.aerbot.tools import make_award_scenario, run_award_scenario
    cid = condsvc.add_condition("only suppliers from India", source="test")["condition_id"]
    try:
        out = run_award_scenario(make_award_scenario("SCN-T", "t"))
        assert "SUP-003" not in out["awards_by_supplier"], \
            "Singapore supplier must be filtered by region rule"
        assert any("India" in s for s in out["conditions_applied"])
    finally:
        condsvc.delete_condition(cid)
    assert run_award_scenario(make_award_scenario("SCN-T", "t"))["line_count"] == 28


# --------------------------------------------------------------------- brief
def test_full_brief_covers_all_lines(conn):
    import json
    from app.aerbot.tools import get_full_brief
    b = get_full_brief()
    assert len(b["lines"]) == 28, "every awarded line exactly once"
    assert all(l["cheapest"] and l["cheapest_cost"] for l in b["lines"])
    assert "80 eligible" in b["summary"]
    l7 = next(l for l in b["lines"] if l["line"] == 7)
    assert (l7["cheapest"], l7["cheapest_cost"]) == ("SUP-001", 11.28125)
    assert len(json.dumps(b)) < 20000, "brief must stay compact for bulk use"


# ------------------------------------------------------- provider fallback
def test_chain_order_prefers_cloud(monkeypatch):
    from app.aerbot import llm
    monkeypatch.setenv("KTQ_PROVIDER_MODE", "auto")
    monkeypatch.setenv("KTQ_LLM_URL", "https://api.example.com/v1")
    monkeypatch.setenv("KTQ_LLM_KEY", "k")
    monkeypatch.setenv("KTQ_LLM_MODEL", "m-cloud")
    monkeypatch.setattr(llm, "_ollama_reachable", lambda base: True)
    chain = llm.resolve_chain()
    assert chain[0][0] == "openai" and chain[-1][0] == "ollama"


def test_chain_falls_back_on_429(monkeypatch):
    import httpx
    from app.aerbot import llm
    calls = []

    def fake_dispatch(provider, cfg, messages, system, tools):
        calls.append(provider)
        if provider == "openai":
            req = httpx.Request("POST", "http://x")
            raise httpx.HTTPStatusError("busy", request=req,
                                        response=httpx.Response(429, request=req))
        return llm.ChatResponse(content="ok-local")

    monkeypatch.setattr(llm, "_dispatch", fake_dispatch)
    monkeypatch.setattr(llm, "resolve_chain",
                        lambda: [("openai", {"model": "m"}),
                                 ("ollama", {"model": "q"})])
    r = llm.client_chat("s", [], [])
    assert r.content == "ok-local" and calls == ["openai", "ollama"]
    assert llm.LAST_PROVIDER["name"] == "ollama:q"


def test_chain_401_fails_fast(monkeypatch):
    import httpx
    from app.aerbot import llm
    calls = []

    def fake_dispatch(provider, cfg, messages, system, tools):
        calls.append(provider)
        req = httpx.Request("POST", "http://x")
        raise httpx.HTTPStatusError("bad key", request=req,
                                    response=httpx.Response(401, request=req))

    monkeypatch.setattr(llm, "_dispatch", fake_dispatch)
    monkeypatch.setattr(llm, "resolve_chain",
                        lambda: [("openai", {"model": "m"}),
                                 ("ollama", {"model": "q"})])
    try:
        llm.client_chat("s", [], [])
        assert False, "401 must surface, not silently fall back"
    except httpx.HTTPStatusError:
        pass
    assert calls == ["openai"]


def test_chain_empty_means_unconfigured(monkeypatch):
    from app.aerbot import llm
    monkeypatch.setattr(llm, "resolve_chain", lambda: [])
    try:
        llm.client_chat("s", [], [])
        assert False
    except llm.ModelNotConfigured:
        pass


# ------------------------------------------------------- provider settings
def test_mode_offline_excludes_cloud(monkeypatch):
    from app.aerbot import llm
    monkeypatch.setenv("KTQ_PROVIDER_MODE", "offline")
    monkeypatch.setenv("KTQ_LLM_URL", "https://cloud.example/v1")
    monkeypatch.setenv("KTQ_LLM_KEY", "k")
    monkeypatch.setattr(llm, "_ollama_reachable", lambda base: True)
    assert [p for p, _ in llm.resolve_chain()] == ["ollama"]


def test_mode_keyless_inserts_demo_cloud(monkeypatch):
    from app.aerbot import llm
    monkeypatch.setenv("KTQ_PROVIDER_MODE", "keyless")
    for k in ("KTQ_LLM_URL", "KTQ_LLM_KEY", "OPENAI_API_KEY", "ANTHROPIC_API_KEY"):
        monkeypatch.delenv(k, raising=False)
    monkeypatch.setattr(llm, "_ollama_reachable", lambda base: True)
    chain = llm.resolve_chain()
    assert chain[0][0] == "openai" and "pollinations" in chain[0][1]["url"]
    assert chain[0][1]["key"] is None and chain[-1][0] == "ollama"


def test_mode_groq_needs_valid_key(monkeypatch, tmp_path):
    from app.aerbot import llm
    monkeypatch.setattr(llm, "validate_groq_key", lambda k: k == "gsk-good")
    monkeypatch.setattr(llm, "_env_path", lambda: tmp_path / ".env")
    (tmp_path / ".env").write_text("# test\nOLLAMA_HOST=http://x\n")
    try:
        llm.set_provider_settings("groq", "gsk-bad")
        assert False, "bad key must be rejected"
    except ValueError:
        pass
    s = llm.set_provider_settings("groq", "gsk-good")
    assert s["mode"] == "groq" and s["has_groq_key"] is True
    text = (tmp_path / ".env").read_text()
    assert "KTQ_LLM_URL=https://api.groq.com/openai/v1" in text
    assert "gsk-good" in text and "# test" in text
    for k in ("KTQ_PROVIDER_MODE", "KTQ_LLM_URL", "KTQ_LLM_KEY", "KTQ_LLM_MODEL",
              "OLLAMA_HOST"):
        monkeypatch.delenv(k, raising=False)


def test_settings_never_leak_key(monkeypatch):
    import inspect
    from app.aerbot import llm
    src = inspect.getsource(llm.provider_settings)
    assert "KTQ_LLM_KEY" not in src and "API_KEY" not in src


def test_chain_400_falls_back(monkeypatch):
    import httpx
    from app.aerbot import llm
    calls = []

    def fake_dispatch(provider, cfg, messages, system, tools):
        calls.append(provider)
        if provider == "openai":
            req = httpx.Request("POST", "http://x")
            raise httpx.HTTPStatusError("quirky backend", request=req,
                                        response=httpx.Response(400, request=req))
        return llm.ChatResponse(content="ok-local")

    monkeypatch.setattr(llm, "_dispatch", fake_dispatch)
    monkeypatch.setattr(llm, "resolve_chain",
                        lambda: [("openai", {"model": "m"}),
                                 ("ollama", {"model": "q"})])
    assert llm.client_chat("s", [], []).content == "ok-local"
    assert calls == ["openai", "ollama"]


def test_tool_schema_has_no_nones():
    import json
    from app.aerbot import agent as ag
    from app.aerbot.llm import _openai_tool_schema
    schema = json.loads(json.dumps(_openai_tool_schema(ag.TOOLS)))

    def walk(node):
        if isinstance(node, dict):
            assert None not in node.values()
            for v in node.values():
                walk(v)
        elif isinstance(node, list):
            for v in node:
                walk(v)

    walk(schema)


def test_settings_stats_and_ocr_lang_roundtrip(conn, tmp_path, monkeypatch):
    from app import main as appmain
    from app.aerbot import llm
    s = appmain.settings_stats()
    assert s["offers"] == 131 and s["eligible"] == 80 and s["lines_awarded"] == 28
    assert set(s) >= {"uploads", "conditions", "upload_files_on_disk"}
    monkeypatch.setattr(llm, "_env_path", lambda: tmp_path / ".env")
    (tmp_path / ".env").write_text("A=1\n")
    out = llm.set_provider_settings("offline", ocr_lang="eng+hin")
    assert out["ocr_lang"] == "eng+hin"
    text = (tmp_path / ".env").read_text()
    assert "KTQ_OCR_LANG=eng+hin" in text and "A=1" in text
    for k in ("KTQ_PROVIDER_MODE", "KTQ_OCR_LANG"):
        monkeypatch.delenv(k, raising=False)


def test_chitchat_skips_tools_and_summaries(scripted):
    import json
    from app.aerbot import answer
    from app.aerbot.agent import _is_chitchat
    from app.aerbot.llm import ChatResponse
    assert _is_chitchat("Hello!") and _is_chitchat("who are you?")
    assert _is_chitchat("thanks") and not _is_chitchat("how are your prices?")
    assert not _is_chitchat("Give me a summary of the RFx")
    seen = {}

    def chat_fn(system, messages, tools):
        seen["tools"] = tools
        return ChatResponse(content=json.dumps({
            "intent": "query", "title": "AerBot",
            "narrative": "Hello! I'm AerBot, your procurement copilot.",
            "table": [], "assumptions": [], "exclusions": [],
            "uncertainties": [], "evidence": []}))

    a = answer("Hello!", chat_fn=chat_fn)
    assert seen["tools"] == [], "greetings must offer no tools"
    assert "131 offers" not in a["narrative"]
    assert "AerBot" in a["narrative"]

    def mute(system, messages, tools):
        return ChatResponse(content=None)

    b = answer("hi", chat_fn=mute)
    assert "AerBot" in b["narrative"], "empty model turn still gets a greeting"


def test_docx_extracts_paragraphs_and_tables(tmp_path):
    from docx import Document
    from app.services.extraction import extract_docx
    p = str(tmp_path / "q.docx")
    doc = Document()
    doc.add_paragraph("Quotation for your RFx")
    doc.add_paragraph("D-003 Qty: 500 box Rs. 120 10d valid 2026-11-30")
    tbl = doc.add_table(rows=2, cols=2)
    tbl.cell(0, 0).text = "Code"
    tbl.cell(0, 1).text = "Price"
    tbl.cell(1, 0).text = "D-005"
    tbl.cell(1, 1).text = "CARD 200 INR 88.5 12d"
    doc.save(p)
    rows = extract_docx(p)
    assert rows and rows[0]["quantity"] == 500.0 and rows[0]["currency"] == "INR"


def test_rfx_draft_rejects_unknown_sku(conn):
    from app.services import drafts
    try:
        drafts.create_draft("t", lines_text="Item-006 x 5")
        ok = True
    except ValueError:
        ok = False
    assert ok
    try:
        drafts.create_draft("t", lines_text="Item-999 x 5")
        assert False, "unknown SKUs need confirmation, never silent invention"
    except drafts.NeedConfirm as e:
        assert e.unknown[0]["sku"] == "Item-999" and e.parsed_ok == 1
    ds = drafts.list_drafts()
    assert any(d["name"] == "t" for d in ds)
    conn.execute("DELETE FROM rfx_draft_line WHERE draft_id IN "
                 "(SELECT draft_id FROM rfx_draft WHERE name='t')")
    conn.execute("DELETE FROM rfx_draft WHERE name='t'")
    conn.commit()


def test_rfx_draft_confirm_creates_sku(conn):
    from app.services import drafts
    r = drafts.create_draft(
        "tc", lines_text="Item-006 x 5\nItem-888 x 10",
        new_skus=[{"sku_id": "Item-888", "name": "Test Widget", "uom": "BOX"}],
        source="test")
    assert r["skus_created"] == ["Item-888"]
    d = drafts.get_draft(r["draft_id"])
    assert len(d["lines"]) == 2
    assert d["new_skus"][0]["sku_name"] == "Test Widget"
    origins = {l["sku_id"]: l["sku_origin"] for l in d["lines"]}
    assert origins["Item-888"] == "draft" and origins["Item-006"] == "master"
    # master catalogue untouched
    assert conn.execute("SELECT COUNT(*) FROM sku").fetchone()[0] == 50
    try:
        drafts.create_draft("bad", lines_text="Item-889 x 1",
                            new_skus=[{"sku_id": "Item-889", "name": "", "uom": "BOX"}])
        assert False, "nameless SKUs must be rejected"
    except ValueError:
        pass
    conn.execute("DELETE FROM rfx_draft_line WHERE draft_id=?", (r["draft_id"],))
    conn.execute("DELETE FROM rfx_draft_sku WHERE draft_id=?", (r["draft_id"],))
    conn.execute("DELETE FROM rfx_draft WHERE draft_id=?", (r["draft_id"],))
    conn.commit()


def test_pipeline_validate_reports(conn):
    from app import pipeline
    res = pipeline.validate()
    assert isinstance(res, list) and len(res) >= 8
    assert all("check" in r and "ok" in r for r in res)


def test_draft_validation_and_delete(conn):
    from app.services import drafts
    for bad, why in [("Item-006", "quantity"), ("Item-006 x 5", "deadline")]:
        try:
            drafts.create_draft("v", lines_text=bad,
                                deadline="someday" if why == "deadline" else "")
            assert False, f"{why} must be validated"
        except ValueError as e:
            assert why in str(e)
    r = drafts.create_draft("vok", lines_text="Item-006 x 5",
                            deadline="2026-12-31", source="test")
    assert drafts.delete_draft(r["draft_id"]) is True
    assert drafts.delete_draft(r["draft_id"]) is False
    assert drafts.get_draft(r["draft_id"]) is None


def test_agent_history_flows_into_messages():
    import json
    from app.aerbot import agent as ag
    from app.aerbot.llm import ChatResponse, ToolCall
    seen = {}
    state = {"n": 0}

    def chat_fn(system, messages, tools):
        state["n"] += 1
        if state["n"] == 1:
            seen["roles"] = [m["role"] for m in messages]
            return ChatResponse(tool_calls=[ToolCall(id="t1", name="get_rfx_summary",
                                                     arguments={})])
        return ChatResponse(content=json.dumps({
            "intent": "query", "title": "t", "narrative": "done",
            "table": [], "assumptions": [], "exclusions": [],
            "uncertainties": [], "evidence": []}))

    ag.run("Item-999 is Test Widget, UOM BOX",
           chat_fn=chat_fn,
           history=[{"role": "user", "content": "draft with Item-999"},
                    {"role": "assistant", "content": "which name and UOM?"}])
    assert seen["roles"] == ["user", "assistant", "user"]


def test_session_memory_roundtrip(monkeypatch):
    import json
    from app import main as appmain
    from app.aerbot import llm
    calls = []

    def fake_answer(question, history=None, chat_fn=None):
        calls.append({"q": question, "hist": list(history or [])})
        return {"question": question, "intent": "query", "title": "t",
                "narrative": f"heard: {question}", "table": [],
                "assumptions": [], "exclusions": [], "uncertainties": [],
                "evidence": [], "provider": "test", "markdown": ""}

    monkeypatch.setattr(appmain, "aerbot_answer", fake_answer)
    q1 = appmain.Question(question="draft with Item-999 please", session_id="s-test-1")
    q2 = appmain.Question(question="Item-999 is Test Widget, UOM BOX", session_id="s-test-1")
    appmain.aerbot(q1)
    appmain.aerbot(q2)
    assert len(calls) == 2
    assert calls[0]["hist"] == []
    assert [m["content"] for m in calls[1]["hist"]] == [
        "draft with Item-999 please", "heard: draft with Item-999 please"]
    appmain.aerbot_forget(appmain.SessionOnly(session_id="s-test-1"))
    assert "s-test-1" not in appmain.SESSIONS


def test_draft_freetext_name_proposes_and_confirms(conn):
    from app.services import drafts
    try:
        drafts.create_draft("ft", lines_text="Bla x 2000")
        assert False, "free-text items need confirmation"
    except drafts.NeedConfirm as e:
        assert e.unknown[0]["name"] == "Bla"
        assert e.unknown[0]["suggested_code"] == "Item-051"
    r = drafts.create_draft(
        "ft", lines_text="Bla x 2000",
        new_skus=[{"sku_id": "", "name": "Bla Widget", "uom": "BOX", "line_no": 1}],
        source="test")
    assert r["skus_created"] == ["Item-051"]
    d = drafts.get_draft(r["draft_id"])
    assert d["lines"][0]["sku_id"] == "Item-051"
    assert d["lines"][0]["quantity"] == 2000.0
    assert conn.execute("SELECT COUNT(*) FROM sku").fetchone()[0] == 50
    conn.execute("DELETE FROM rfx_draft_line WHERE draft_id=?", (r["draft_id"],))
    conn.execute("DELETE FROM rfx_draft_sku WHERE draft_id=?", (r["draft_id"],))
    conn.execute("DELETE FROM rfx_draft WHERE draft_id=?", (r["draft_id"],))
    conn.commit()


def test_model_inventory_lists_and_flags(monkeypatch):
    import httpx
    from app.aerbot import llm

    class FakeResp:
        status_code = 200

        def json(self):
            return {"models": [{"name": "qwen2.5:7b-instruct", "size": 4680000000,
                                "modified_at": "2026-09-18T00:00:00Z"}]}

    monkeypatch.setattr(httpx, "get", lambda *a, **k: FakeResp())
    inv = llm.ollama_inventory()
    assert inv["reachable"] is True and inv["installed"] is True
    assert inv["models"][0]["size_gb"] == 4.68

    def boom(*a, **k):
        raise httpx.ConnectError("down")

    monkeypatch.setattr(httpx, "get", boom)
    inv = llm.ollama_inventory()
    assert inv["reachable"] is False and inv["installed"] is False


def test_model_pull_rejects_bad_names():
    from app.aerbot import llm
    for bad in ["../../etc/passwd", "a" * 200, "model; rm -rf /", "m$(x)"]:
        try:
            llm.start_model_pull(bad)
            assert False, f"must reject {bad!r}"
        except ValueError:
            pass


def test_pull_status_initial_shape():
    from app.aerbot import llm
    s = llm.pull_status()
    assert set(s) >= {"running", "model", "done", "error", "log"}


def test_inventory_reports_models_dir(monkeypatch):
    import httpx
    from app.aerbot import llm

    class FakeResp:
        status_code = 200

        def json(self):
            return {"models": []}

    monkeypatch.setattr(httpx, "get", lambda *a, **k: FakeResp())
    inv = llm.ollama_inventory()
    assert "models_dir" in inv and inv["models_dir"]
